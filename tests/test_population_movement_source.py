"""Tests for src/fetchers/population_movement.py.

Fixture workbooks are built in-memory with openpyxl rather than committing a
slice of the real 3.4 MB file. The real header layout (rows 2-4, merged
group-label cells at their own origin column) is reproduced exactly --
verified directly against the real file, 2026-09-23 -- so a header-lookup
bug cannot hide behind a too-convenient fixture shape. Values for Namur
(92094), Liège (62063) and Aartselaar (11001) 2025 below are hand-copied
from the real file.
"""

import io
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.fetchers.population_movement import (
    BIRTHS,
    DEATHS,
    INTERNAL_MIGRATION_IN,
    INTERNAL_MIGRATION_NET,
    INTERNAL_MIGRATION_OUT,
    INTERNATIONAL_MIGRATION_NET,
    PopulationMovementLinkNotFoundError,
    PopulationMovementSchemaError,
    PopulationMovementSource,
    discover_xlsx_url,
)

# Row 2 (group labels, written only at each merge's own origin column),
# row 3 (ENTREES/SORTIES/SOLDE sub-groups), row 4 (leaf labels) -- all 25
# columns A-Y, exactly as the real file has them.
ROW2 = [
    "CODE INS",
    "LIEU DE RESIDENCE",
    "POPULATION AU 1 JANVIER",
    "MOUVEMENT NATUREL",
    None,
    None,
    "MOUVEMENT MIGRATOIRE INTERNE",
    None,
    None,
    "MOUVEMENT MIGRATOIRE INTERNATIONAL",
    None,
    None,
    None,
    None,
    None,
    None,
    "TOTAL DES MOUVEMENTS DE POPULATION",
    None,
    None,
    "AJUSTEMENT STATISTIQUE",
    "CHANGEMENTS DE NATIONALITE",
    None,
    "POPULATION AU 31 DECEMBRE (SOIT AU 1/1 DE L'ANNEE SUIVANTE)",
    "ACCROISSEMENT",
    None,
]
ROW3 = [None] * 9 + [
    "ENTREES",
    None,
    None,
    "SORTIES",
    None,
    None,
    "SOLDE",
    None,
    None,
    None,
    None,
    None,
    None,
]
ROW4 = [
    None,
    None,
    None,
    "NAISSANCES",
    "DECES",
    "SOLDE",
    "ENTREES",
    "SORTIES",
    "SOLDE",
    "IMMIGRATIONS",
    "CHANGE-MENTS REGISTRE (ENTREES)",
    "REINSCRITS AYANT ETE RAYES",
    "EMIGRATIONS",
    "CHANGE-MENTS REGISTRE (SORTIES)",
    "RAYES D'OFFICE",
    None,
    "ENTREES",
    "SORTIES",
    "SOLDE",
    None,
    "BELGE EN NON-BELGE",
    "NON-BELGE EN BELGE",
    None,
    "TOTAL",
    "INDICE",
]


def _data_row(
    code: str,
    naissances: float,
    deces: float,
    internal_solde: float,
    international_solde: float,
    internal_entrees: float | None = None,
    internal_sorties: float | None = None,
) -> list:
    row = [None] * 25
    row[0] = code
    row[1] = f"Commune {code}"
    row[2] = 10000
    row[3] = naissances
    row[4] = deces
    row[5] = naissances - deces
    # Default ENTREES/SORTIES to values consistent with the given SOLDE when
    # not supplied explicitly, so existing callers (which only pass SOLDE)
    # keep working without every fixture needing updating.
    if internal_entrees is None:
        internal_entrees = max(internal_solde, 0)
    if internal_sorties is None:
        internal_sorties = internal_entrees - internal_solde
    row[6] = internal_entrees
    row[7] = internal_sorties
    row[8] = internal_solde
    row[15] = international_solde
    return row


def _make_workbook(sheets: dict[str, list[list]]) -> bytes:
    """`sheets`: {sheet_name: [data_row, ...]} -- rows built by _data_row()."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, data_rows in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(["Mouvement de la population en " + name] + [None] * 24)
        ws.append(ROW2)
        ws.append(ROW3)
        ws.append(ROW4)
        for row in data_rows:
            ws.append(row)
        # The real file's merges -- reproduced for realism; read_only mode
        # (what the adapter uses) reads only the origin cell of a merge, so
        # this is not load-bearing for the adapter itself, only fidelity.
        ws.merge_cells("D2:F2")
        ws.merge_cells("G2:I2")
        ws.merge_cells("J2:P2")
        ws.merge_cells("J3:L3")
        ws.merge_cells("M3:O3")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ENTREES/SORTIES (columns G/H) hand-copied from the real 2025 sheet,
# 2026-09-24: Namur 5258/5558 (5258-5558=-300, matches the existing SOLDE
# column I exactly), Liège 10411/11704 (10411-11704=-1293, matches).
REAL_2025_NAMUR = _data_row("92094", 1036, 1118, -300, 826, 5258, 5558)
REAL_2025_LIEGE = _data_row("62063", 2016, 2024, -1293, 2135, 10411, 11704)
REAL_2025_AARTSELAAR = _data_row("11001", 131, 170, 34, 0)


# --- link discovery ----------------------------------------------------------


def test_discover_xlsx_url_finds_the_real_link_shape():
    html = (
        '<a href="https://statbel.fgov.be/sites/default/files/files/documents/bevolking/'
        '5.2%20Loop%20van%20de%20bevolking/pop1992-mov_fr.xlsx">Download</a>'
    )
    url = discover_xlsx_url(html)
    assert url.endswith("pop1992-mov_fr.xlsx")


def test_discover_xlsx_url_raises_when_not_found():
    with pytest.raises(PopulationMovementLinkNotFoundError):
        discover_xlsx_url("<html><body>no link here</body></html>")


def test_discover_xlsx_url_raises_on_multiple_candidates():
    html = (
        '<a href="https://statbel.fgov.be/a/pop1992-mov_fr.xlsx">A</a>'
        '<a href="https://statbel.fgov.be/b/pop1992-mov_nl.xlsx">B</a>'
    )
    # Both hrefs contain "pop1992-mov" and end in .xlsx -- two distinct
    # candidates, refused rather than guessed at.
    with pytest.raises(PopulationMovementLinkNotFoundError):
        discover_xlsx_url(html)


# --- parsing -------------------------------------------------------------


def test_parses_hand_computed_2025_values():
    raw = _make_workbook({"2025": [REAL_2025_NAMUR, REAL_2025_LIEGE, REAL_2025_AARTSELAAR]})
    rows = PopulationMovementSource()._parse(raw)

    by_key = {(r["indicator_id"], r["geo_id"], r["period"]): r["value"] for r in rows}
    assert by_key[(BIRTHS, "92094", "2025")] == 1036.0
    assert by_key[(DEATHS, "92094", "2025")] == 1118.0
    assert by_key[(INTERNAL_MIGRATION_NET, "92094", "2025")] == -300.0
    assert by_key[(INTERNATIONAL_MIGRATION_NET, "92094", "2025")] == 826.0
    assert by_key[(INTERNAL_MIGRATION_IN, "92094", "2025")] == 5258.0
    assert by_key[(INTERNAL_MIGRATION_OUT, "92094", "2025")] == 5558.0
    # IN - OUT reproduces the publisher's own SOLDE exactly.
    assert (
        by_key[(INTERNAL_MIGRATION_IN, "92094", "2025")]
        - by_key[(INTERNAL_MIGRATION_OUT, "92094", "2025")]
        == by_key[(INTERNAL_MIGRATION_NET, "92094", "2025")]
    )

    assert by_key[(BIRTHS, "62063", "2025")] == 2016.0
    assert by_key[(DEATHS, "62063", "2025")] == 2024.0
    assert by_key[(INTERNAL_MIGRATION_NET, "62063", "2025")] == -1293.0
    assert by_key[(INTERNATIONAL_MIGRATION_NET, "62063", "2025")] == 2135.0
    assert by_key[(INTERNAL_MIGRATION_IN, "62063", "2025")] == 10411.0
    assert by_key[(INTERNAL_MIGRATION_OUT, "62063", "2025")] == 11704.0
    assert (
        by_key[(INTERNAL_MIGRATION_IN, "62063", "2025")]
        - by_key[(INTERNAL_MIGRATION_OUT, "62063", "2025")]
        == by_key[(INTERNAL_MIGRATION_NET, "62063", "2025")]
    )

    assert by_key[(BIRTHS, "11001", "2025")] == 131.0
    assert by_key[(DEATHS, "11001", "2025")] == 170.0
    assert by_key[(INTERNAL_MIGRATION_NET, "11001", "2025")] == 34.0


def test_every_row_carries_the_base_contract_keys():
    raw = _make_workbook({"2025": [REAL_2025_NAMUR]})
    rows = PopulationMovementSource()._parse(raw)
    assert rows, "fixture must produce at least one row"
    for row in rows:
        assert {"geo_id", "period", "value", "status"} <= set(row.keys())
        assert isinstance(row["geo_id"], str)
        assert isinstance(row["period"], str)
        assert isinstance(row["value"], float)
        assert row["status"] == "final"
        assert row["indicator_id"] in {
            BIRTHS,
            DEATHS,
            INTERNAL_MIGRATION_NET,
            INTERNATIONAL_MIGRATION_NET,
            INTERNAL_MIGRATION_IN,
            INTERNAL_MIGRATION_OUT,
        }


def test_six_indicators_emitted_per_data_row():
    raw = _make_workbook({"2025": [REAL_2025_NAMUR]})
    rows = PopulationMovementSource()._parse(raw)
    assert len(rows) == 6
    assert {r["indicator_id"] for r in rows} == {
        BIRTHS,
        DEATHS,
        INTERNAL_MIGRATION_NET,
        INTERNATIONAL_MIGRATION_NET,
        INTERNAL_MIGRATION_IN,
        INTERNAL_MIGRATION_OUT,
    }


def test_multiple_sheets_all_parsed():
    raw = _make_workbook(
        {"2025": [REAL_2025_NAMUR], "1992": [_data_row("92094", 900, 1000, -50, 20)]}
    )
    rows = PopulationMovementSource()._parse(raw)
    periods = {r["period"] for r in rows}
    assert periods == {"2025", "1992"}


def test_footer_and_blank_rows_are_skipped_not_errored():
    raw = _make_workbook({"2025": [REAL_2025_NAMUR]})
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    ws = wb["2025"]
    ws.append(["Source: Statbel (Direction générale Statistique)"] + [None] * 24)
    ws.append([None] * 25)
    buf = io.BytesIO()
    wb.save(buf)
    rows = PopulationMovementSource()._parse(buf.getvalue())
    assert len(rows) == 6  # only Namur's row produced data


# --- schema refusals (CLAUDE.md rule 13) -------------------------------------


def test_no_year_sheet_refuses():
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(PopulationMovementSchemaError, match="year"):
        PopulationMovementSource()._parse(buf.getvalue())


def test_code_ins_header_moved_refuses():
    row2 = list(ROW2)
    row2[0] = "SOMETHING ELSE"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("2025")
    ws.append(["title"] + [None] * 24)
    ws.append(row2)
    ws.append(ROW3)
    ws.append(ROW4)
    ws.append(REAL_2025_NAMUR)
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(PopulationMovementSchemaError, match="CODE INS"):
        PopulationMovementSource()._parse(buf.getvalue())


def test_naissances_column_header_moved_refuses():
    row4 = list(ROW4)
    row4[3] = "SOMETHING ELSE"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("2025")
    ws.append(["title"] + [None] * 24)
    ws.append(ROW2)
    ws.append(ROW3)
    ws.append(row4)
    ws.append(REAL_2025_NAMUR)
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(PopulationMovementSchemaError, match="NAISSANCES"):
        PopulationMovementSource()._parse(buf.getvalue())


def test_internal_migration_group_label_moved_refuses():
    row2 = list(ROW2)
    row2[6] = "SOMETHING ELSE"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("2025")
    ws.append(["title"] + [None] * 24)
    ws.append(row2)
    ws.append(ROW3)
    ws.append(ROW4)
    ws.append(REAL_2025_NAMUR)
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(PopulationMovementSchemaError, match="MOUVEMENT MIGRATOIRE INTERNE"):
        PopulationMovementSource()._parse(buf.getvalue())


def test_internal_entrees_column_header_moved_refuses():
    row4 = list(ROW4)
    row4[6] = "SOMETHING ELSE"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("2025")
    ws.append(["title"] + [None] * 24)
    ws.append(ROW2)
    ws.append(ROW3)
    ws.append(row4)
    ws.append(REAL_2025_NAMUR)
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(PopulationMovementSchemaError, match="ENTREES"):
        PopulationMovementSource()._parse(buf.getvalue())


def test_internal_sorties_column_header_moved_refuses():
    row4 = list(ROW4)
    row4[7] = "SOMETHING ELSE"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("2025")
    ws.append(["title"] + [None] * 24)
    ws.append(ROW2)
    ws.append(ROW3)
    ws.append(row4)
    ws.append(REAL_2025_NAMUR)
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(PopulationMovementSchemaError, match="SORTIES"):
        PopulationMovementSource()._parse(buf.getvalue())


def test_non_numeric_births_cell_refuses():
    row = list(REAL_2025_NAMUR)
    row[3] = "non disponible"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("2025")
    ws.append(["title"] + [None] * 24)
    ws.append(ROW2)
    ws.append(ROW3)
    ws.append(ROW4)
    ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(PopulationMovementSchemaError, match="not numeric"):
        PopulationMovementSource()._parse(buf.getvalue())
